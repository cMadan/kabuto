#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook

from kabuto_common import (
    ensure_sheet_headers,
    excel_row_to_question,
    parse_chapter_section_numbers,
    validate_questions,
)


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Validate curated Excel MCQ bank (default), with optional summary reporting"
    )
    p.add_argument("input_xlsx", help="Path to Excel workbook")
    p.add_argument("--sheet", default=None, help="Worksheet name (default: active sheet)")
    p.add_argument(
        "--warnings-as-errors",
        action="store_true",
        help="Return nonzero exit code if warnings exist (validation modes only)",
    )

    p.add_argument("--summary", action="store_true", help="Print bank summary in addition to validation output")
    p.add_argument(
        "--summary-only",
        action="store_true",
        help="Print summary only (does not fail on validation errors)",
    )
    p.add_argument("--json-out", default=None, help="Write summary JSON to this path (summary modes)")
    p.add_argument("--show-sections", type=int, default=0, help="Show top N chapter_section rows (0 = all)")
    p.add_argument(
        "--show-warning-samples",
        type=int,
        default=10,
        help="Show up to N warning lines in summary (default: 10)",
    )
    p.add_argument(
        "--show-error-samples",
        type=int,
        default=10,
        help="Show up to N error lines in summary (default: 10)",
    )
    return p


def _section_sort_key(s: str) -> Tuple[int, int, str]:
    nums = parse_chapter_section_numbers(s)
    if nums is None:
        return (10**9, 10**9, s.lower())
    return (nums[0], nums[1], s.lower())


def _warning_category(msg: str) -> str:
    m = msg.lower()
    if "duplicate option text" in m:
        return "duplicate_option_text"
    if "letter-referential" in m:
        return "letter_referential_with_random_shuffle"
    if "missing chapter_section" in m:
        return "missing_chapter_section"
    return "other"


def _error_category(msg: str) -> str:
    m = msg.lower()
    if m.startswith("row "):
        return "structural_row_parse"
    if "duplicate id" in m:
        return "duplicate_id"
    if "missing id" in m:
        return "missing_id"
    if "missing stem" in m:
        return "missing_stem"
    if "fewer than 2 options" in m:
        return "too_few_options"
    if "more than 5 options" in m:
        return "too_many_options"
    if "gapped options" in m:
        return "gapped_options"
    if "invalid answer" in m:
        return "invalid_answer"
    if "points to missing option" in m:
        return "answer_points_to_missing_option"
    if "invalid shuffle_mode" in m:
        return "invalid_shuffle_mode"
    if "points" in m and "invalid" in m:
        return "invalid_points"
    if "active" in m and ("0 or 1" in m or "invalid active" in m):
        return "invalid_active"
    return "other"


def _load_questions(ws: Any):
    questions = []
    structural_errors: List[str] = []
    nonempty_rows = 0
    empty_rows = 0
    for r in range(2, ws.max_row + 1):
        row_values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if all(v is None or str(v).strip() == "" for v in row_values):
            empty_rows += 1
            continue
        nonempty_rows += 1
        try:
            q = excel_row_to_question(ws, r)
        except Exception as exc:
            structural_errors.append(f"Row {r}: {exc}")
            continue
        if q is not None:
            questions.append(q)
    return questions, structural_errors, nonempty_rows, empty_rows


def _build_summary(
    input_xlsx: str,
    ws: Any,
    questions: List[Any],
    errors: List[str],
    warnings: List[str],
    nonempty_rows: int,
    empty_rows: int,
    show_sections: int,
    show_error_samples: int,
    show_warning_samples: int,
) -> Dict[str, Any]:
    active_qs = [q for q in questions if int(getattr(q, "active", 1)) == 1]
    inactive_qs = [q for q in questions if int(getattr(q, "active", 1)) != 1]

    shuffle_counts_total = Counter((q.shuffle_mode or "").strip().lower() or "random" for q in questions)
    shuffle_counts_active = Counter((q.shuffle_mode or "").strip().lower() or "random" for q in active_qs)

    explanation_missing_total = sum(1 for q in questions if not (q.explanation and str(q.explanation).strip()))
    explanation_missing_active = sum(1 for q in active_qs if not (q.explanation and str(q.explanation).strip()))

    section_rows: Dict[str, Dict[str, int]] = defaultdict(lambda: {"total": 0, "active": 0, "inactive": 0})
    for q in questions:
        sec = (q.chapter_section or "").strip() or "<missing>"
        section_rows[sec]["total"] += 1
        if int(getattr(q, "active", 1)) == 1:
            section_rows[sec]["active"] += 1
        else:
            section_rows[sec]["inactive"] += 1

    warning_categories = Counter(_warning_category(w) for w in warnings)
    error_categories = Counter(_error_category(e) for e in errors)

    sections_sorted = sorted(section_rows.keys(), key=_section_sort_key)
    top_n = show_sections if show_sections and show_sections > 0 else len(sections_sorted)
    sections_preview = [
        {
            "chapter_section": s,
            "total": section_rows[s]["total"],
            "active": section_rows[s]["active"],
            "inactive": section_rows[s]["inactive"],
        }
        for s in sections_sorted[:top_n]
    ]

    return {
        "file": str(Path(input_xlsx)),
        "sheet": ws.title,
        "worksheet": {
            "max_row": ws.max_row,
            "nonempty_data_rows": nonempty_rows,
            "empty_data_rows": empty_rows,
            "parsed_questions": len(questions),
        },
        "questions": {
            "total": len(questions),
            "active": len(active_qs),
            "inactive": len(inactive_qs),
            "missing_explanation_total": explanation_missing_total,
            "missing_explanation_active": explanation_missing_active,
            "shuffle_mode_total": dict(sorted(shuffle_counts_total.items())),
            "shuffle_mode_active": dict(sorted(shuffle_counts_active.items())),
        },
        "validation": {
            "errors": len(errors),
            "warnings": len(warnings),
            "error_categories": dict(error_categories),
            "warning_categories": dict(warning_categories),
            "error_samples": errors[: max(0, show_error_samples)],
            "warning_samples": warnings[: max(0, show_warning_samples)],
        },
        "chapter_section_counts": sections_preview,
        "chapter_section_counts_total_distinct": len(section_rows),
    }


def _print_validation(file_path: str, ws: Any, questions: List[Any], errors: List[str], warnings: List[str]) -> None:
    print(f"File: {Path(file_path)}")
    print(f"Sheet: {ws.title}")
    print(f"Rows parsed: {len(questions)}")
    print(f"Errors: {len(errors)}")
    print(f"Warnings: {len(warnings)}")
    if errors:
        print("\nERRORS")
        for e in errors:
            print(f" - {e}")
    if warnings:
        print("\nWARNINGS")
        for w in warnings:
            print(f" - {w}")


def _print_summary(summary: Dict[str, Any]) -> None:
    print(f"File: {summary['file']}")
    print(f"Sheet: {summary['sheet']}")
    print(f"Rows (non-empty data rows): {summary['worksheet']['nonempty_data_rows']}")
    print(f"Questions parsed: {summary['worksheet']['parsed_questions']}")
    print(f"Active: {summary['questions']['active']} | Inactive: {summary['questions']['inactive']}")
    print(
        "Missing explanation: "
        f"{summary['questions']['missing_explanation_total']} total | "
        f"{summary['questions']['missing_explanation_active']} active"
    )
    print(
        "Validation: "
        f"{summary['validation']['errors']} errors | {summary['validation']['warnings']} warnings"
    )

    sh_total = summary["questions"].get("shuffle_mode_total") or {}
    sh_active = summary["questions"].get("shuffle_mode_active") or {}
    if sh_total:
        print("\nShuffle modes (total):")
        for k, v in sorted(sh_total.items()):
            print(f" - {k}: {v}")
    if sh_active:
        print("Shuffle modes (active):")
        for k, v in sorted(sh_active.items()):
            print(f" - {k}: {v}")

    print("\nChapter-section counts (total/active/inactive):")
    for row in summary["chapter_section_counts"]:
        print(f" - {row['chapter_section']}: {row['total']} / {row['active']} / {row['inactive']}")
    omitted = summary["chapter_section_counts_total_distinct"] - len(summary["chapter_section_counts"])
    if omitted > 0:
        print(f" - ... ({omitted} more sections not shown)")

    ec = summary["validation"].get("error_categories") or {}
    wc = summary["validation"].get("warning_categories") or {}
    if ec:
        print("\nError categories:")
        for k, v in sorted(ec.items(), key=lambda kv: (-kv[1], kv[0])):
            print(f" - {k}: {v}")
    if wc:
        print("\nWarning categories:")
        for k, v in sorted(wc.items(), key=lambda kv: (-kv[1], kv[0])):
            print(f" - {k}: {v}")

    err_samples = summary["validation"].get("error_samples") or []
    warn_samples = summary["validation"].get("warning_samples") or []
    if err_samples:
        print("\nError samples:")
        for e in err_samples:
            print(f" - {e}")
    if warn_samples:
        print("\nWarning samples:")
        for w in warn_samples:
            print(f" - {w}")


def _write_summary_json(summary: Dict[str, Any], json_out: str | None) -> None:
    if not json_out:
        return
    out = Path(json_out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nWrote JSON summary: {out}")


def main() -> int:
    args = build_parser().parse_args()
    if args.summary and args.summary_only:
        raise SystemExit("Use only one of --summary or --summary-only")

    wb = load_workbook(args.input_xlsx)
    ws = wb[args.sheet] if (args.sheet and args.sheet in wb.sheetnames) else wb.active
    ensure_sheet_headers(ws)

    questions, structural_errors, nonempty_rows, empty_rows = _load_questions(ws)
    errors, warnings = validate_questions(questions, require_ids=True)
    errors = structural_errors + errors

    if args.summary_only:
        summary = _build_summary(
            args.input_xlsx,
            ws,
            questions,
            errors,
            warnings,
            nonempty_rows,
            empty_rows,
            args.show_sections,
            args.show_error_samples,
            args.show_warning_samples,
        )
        _print_summary(summary)
        _write_summary_json(summary, args.json_out)
        return 0

    _print_validation(args.input_xlsx, ws, questions, errors, warnings)

    if args.summary:
        print("\n" + "=" * 60 + "\nSUMMARY\n" + "=" * 60)
        summary = _build_summary(
            args.input_xlsx,
            ws,
            questions,
            errors,
            warnings,
            nonempty_rows,
            empty_rows,
            args.show_sections,
            args.show_error_samples,
            args.show_warning_samples,
        )
        _print_summary(summary)
        _write_summary_json(summary, args.json_out)

    if errors:
        return 1
    if warnings and args.warnings_as_errors:
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
