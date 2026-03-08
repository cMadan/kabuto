from __future__ import annotations

import datetime as _dt
import json
import random
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import yaml
from openpyxl import Workbook, load_workbook

EXCEL_COLUMNS = [
    "import_batch",
    "id",
    "chapter_section",
    "stem",
    "opt_a",
    "opt_b",
    "opt_c",
    "opt_d",
    "opt_e",
    "answer",
    "shuffle_mode",
    "points",
    "explanation",
]
OPTION_COLS = ["opt_a", "opt_b", "opt_c", "opt_d", "opt_e"]
OPTION_LETTERS = ["A", "B", "C", "D", "E"]
VALID_SHUFFLE_MODES = {"random", "none"}
DEFAULT_SHUFFLE_MODE = "random"
DEFAULT_POINTS = 1
DEFAULT_SHEET_NAME = "Questions"
DEFAULT_PYEXAM_NAME = "Exam"

LETTER_REFERENCE_PATTERNS = [
    re.compile(r"\bboth\s+[a-e]\s+and\s+[a-e]\b", re.I),
    re.compile(r"\b(?:a|b|c|d|e)\s+and\s+(?:a|b|c|d|e)\b", re.I),
    re.compile(r"\ball\s+of\s+the\s+above\b", re.I),
    re.compile(r"\bnone\s+of\s+the\s+above\b", re.I),
]
CHAPTER_SECTION_RE = re.compile(r"^\s*MTM\s*(\d+)\s*\.\s*(\d+)\s*$", re.I)
ID_RE = re.compile(r"^([a-z0-9_]+)_(\d{3,})$")


@dataclass
class Question:
    import_batch: Optional[str] = None
    qid: Optional[str] = None
    chapter_section: Optional[str] = None
    stem: str = ""
    options: List[str] = None  # type: ignore[assignment]
    answer: str = ""
    shuffle_mode: str = DEFAULT_SHUFFLE_MODE
    points: int = DEFAULT_POINTS
    explanation: Optional[str] = None

    def __post_init__(self) -> None:
        if self.options is None:
            self.options = []

    def correct_index(self) -> int:
        return OPTION_LETTERS.index(self.answer)

    def to_excel_row(self) -> Dict[str, Any]:
        row = {col: "" for col in EXCEL_COLUMNS}
        row["import_batch"] = self.import_batch or ""
        row["id"] = self.qid or ""
        row["chapter_section"] = self.chapter_section or ""
        row["stem"] = self.stem
        for i, opt in enumerate(self.options[:5]):
            row[OPTION_COLS[i]] = opt
        row["answer"] = self.answer
        row["shuffle_mode"] = "" if self.shuffle_mode == DEFAULT_SHUFFLE_MODE else self.shuffle_mode
        row["points"] = "" if self.points == DEFAULT_POINTS else self.points
        row["explanation"] = self.explanation or ""
        return row

    def to_canonical_yaml_obj(self, include_id: bool = True) -> Dict[str, Any]:
        d: Dict[str, Any] = {}
        if include_id and self.qid:
            d["id"] = self.qid
        if self.chapter_section:
            d["chapter_section"] = self.chapter_section
        d["stem"] = self.stem
        d["options"] = list(self.options)
        d["answer"] = self.answer
        if self.shuffle_mode != DEFAULT_SHUFFLE_MODE:
            d["shuffle_mode"] = self.shuffle_mode
        if self.points != DEFAULT_POINTS:
            d["points"] = self.points
        if self.explanation:
            d["explanation"] = self.explanation
        return d


def normalize_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def now_batch_label(prefix: str = "import") -> str:
    return _dt.datetime.now().strftime(f"%Y-%m-%d_{prefix}_%H%M%S")


def normalize_answer(value: Any) -> str:
    s = normalize_text(value).upper()
    if not s:
        raise ValueError("missing answer")
    if s.isdigit():
        i = int(s)
        if 1 <= i <= 5:
            return OPTION_LETTERS[i - 1]
    if s in OPTION_LETTERS:
        return s
    raise ValueError(f"invalid answer '{value}' (expected A-E or 1-5)")


def normalize_shuffle_mode(value: Any, default: str = DEFAULT_SHUFFLE_MODE) -> str:
    s = normalize_text(value).lower()
    if not s:
        return default
    if s not in VALID_SHUFFLE_MODES:
        raise ValueError(f"invalid shuffle_mode '{value}' (expected random|none)")
    return s


def normalize_points(value: Any, default: int = DEFAULT_POINTS) -> int:
    s = normalize_text(value)
    if not s:
        return default
    try:
        x = float(s)
    except Exception as exc:
        raise ValueError(f"invalid points '{value}'") from exc
    if x <= 0 or int(x) != x:
        raise ValueError(f"points must be a positive integer; got '{value}'")
    return int(x)


def normalize_options(options: Any) -> List[str]:
    if not isinstance(options, list):
        raise ValueError("options must be a list")
    cleaned = [normalize_text(x) for x in options]
    cleaned = [x for x in cleaned if x]
    if len(cleaned) < 2:
        raise ValueError("question must have at least 2 non-empty options")
    if len(cleaned) > 5:
        raise ValueError("question may have at most 5 options")
    return cleaned


def parse_chapter_section_prefix(chapter_section: Optional[str]) -> str:
    s = normalize_text(chapter_section)
    if not s:
        return "unassigned"
    m = CHAPTER_SECTION_RE.match(s)
    if not m:
        return "unassigned"
    chapter = int(m.group(1))
    section = int(m.group(2))
    return f"mtm{chapter:02d}_{section:02d}"


def next_id_for_prefix(prefix: str, used_ids: Iterable[str]) -> str:
    max_idx = 0
    for qid in used_ids:
        m = ID_RE.match(qid or "")
        if not m or m.group(1) != prefix:
            continue
        try:
            max_idx = max(max_idx, int(m.group(2)))
        except ValueError:
            pass
    return f"{prefix}_{max_idx + 1:03d}"


def assign_ids(questions: List[Question], existing_ids: Iterable[str]) -> None:
    used = set(existing_ids)
    for q in questions:
        if q.qid:
            used.add(q.qid)
            continue
        prefix = parse_chapter_section_prefix(q.chapter_section)
        q.qid = next_id_for_prefix(prefix, used)
        used.add(q.qid)


def fill_import_batch(questions: List[Question], batch_label: str) -> None:
    for q in questions:
        if not q.import_batch:
            q.import_batch = batch_label


def load_yaml(path: str | Path) -> Any:
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def dump_yaml(obj: Any, path: str | Path) -> None:
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(obj, f, allow_unicode=True, sort_keys=False, width=1000)


def dump_json(obj: Any, path: str | Path) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False)


def load_intake_yaml_questions(yaml_path: str | Path) -> Tuple[Dict[str, Any], List[Question]]:
    data = load_yaml(yaml_path)
    if data is None:
        raise ValueError("YAML file is empty")
    if isinstance(data, list):
        raw_questions = data
        defaults = {}
    elif isinstance(data, dict):
        defaults = data.get("defaults") or {}
        raw_questions = data.get("questions", [])
    else:
        raise ValueError("top-level YAML must be a mapping with 'questions' or a list")
    if not isinstance(raw_questions, list):
        raise ValueError("'questions' must be a list")

    default_shuffle = normalize_shuffle_mode(defaults.get("shuffle_mode", DEFAULT_SHUFFLE_MODE))
    default_points = normalize_points(defaults.get("points", DEFAULT_POINTS))

    questions: List[Question] = []
    for i, raw in enumerate(raw_questions, start=1):
        if not isinstance(raw, dict):
            raise ValueError(f"question {i}: expected mapping")
        if "id" in raw:
            raise ValueError(f"question {i}: intake YAML must not include 'id'")
        stem = normalize_text(raw.get("stem"))
        if not stem:
            raise ValueError(f"question {i}: missing stem")
        options = normalize_options(raw.get("options"))
        answer = normalize_answer(raw.get("answer"))
        if OPTION_LETTERS.index(answer) >= len(options):
            raise ValueError(f"question {i}: answer {answer} points to missing option")
        q = Question(
            chapter_section=normalize_text(raw.get("chapter_section")) or None,
            stem=stem,
            options=options,
            answer=answer,
            shuffle_mode=normalize_shuffle_mode(raw.get("shuffle_mode"), default_shuffle),
            points=normalize_points(raw.get("points"), default_points),
            explanation=normalize_text(raw.get("explanation")) or None,
        )
        questions.append(q)
    return {"shuffle_mode": default_shuffle, "points": default_points}, questions


def create_workbook_with_headers(sheet_name: str = DEFAULT_SHEET_NAME) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(EXCEL_COLUMNS)
    return wb


def ensure_sheet_headers(ws: Any) -> None:
    observed = [normalize_text(ws.cell(row=1, column=i + 1).value) for i in range(len(EXCEL_COLUMNS))]
    if observed != EXCEL_COLUMNS:
        raise ValueError(
            "worksheet header row does not match expected schema.\n"
            f"Expected: {EXCEL_COLUMNS}\nObserved: {observed}"
        )


def load_or_create_workbook(path: str | Path, mode: str, sheet_name: str = DEFAULT_SHEET_NAME):
    mode = mode.lower()
    if mode not in {"new", "append"}:
        raise ValueError("mode must be 'new' or 'append'")
    path = Path(path)
    if mode == "new" or not path.exists():
        wb = create_workbook_with_headers(sheet_name)
        return wb, wb[sheet_name]
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    ensure_sheet_headers(ws)
    return wb, ws


def append_questions_to_sheet(ws: Any, questions: List[Question]) -> None:
    ensure_sheet_headers(ws)
    for q in questions:
        row = q.to_excel_row()
        ws.append([row[col] for col in EXCEL_COLUMNS])


def scan_existing_ids(ws: Any) -> List[str]:
    ensure_sheet_headers(ws)
    id_idx = EXCEL_COLUMNS.index("id") + 1
    out: List[str] = []
    for r in range(2, ws.max_row + 1):
        s = normalize_text(ws.cell(row=r, column=id_idx).value)
        if s:
            out.append(s)
    return out


def excel_row_to_question(ws: Any, row_idx: int) -> Optional[Question]:
    ensure_sheet_headers(ws)
    vals = {col: ws.cell(row=row_idx, column=i + 1).value for i, col in enumerate(EXCEL_COLUMNS)}
    if all(normalize_text(v) == "" for v in vals.values()):
        return None

    options: List[str] = []
    seen_nonblank = False
    seen_gap = False
    for col in OPTION_COLS:
        txt = normalize_text(vals[col])
        if txt:
            if seen_gap:
                raise ValueError("gapped options (blank between non-blank options)")
            seen_nonblank = True
            options.append(txt)
        elif seen_nonblank:
            seen_gap = True

    points_val = vals["points"]
    points = DEFAULT_POINTS if normalize_text(points_val) == "" else int(float(points_val))
    return Question(
        import_batch=normalize_text(vals["import_batch"]) or None,
        qid=normalize_text(vals["id"]) or None,
        chapter_section=normalize_text(vals["chapter_section"]) or None,
        stem=normalize_text(vals["stem"]),
        options=options,
        answer=normalize_text(vals["answer"]).upper(),
        shuffle_mode=normalize_text(vals["shuffle_mode"]).lower() or DEFAULT_SHUFFLE_MODE,
        points=points,
        explanation=normalize_text(vals["explanation"]) or None,
    )


def load_excel_questions(xlsx_path: str | Path, sheet_name: Optional[str] = None) -> List[Question]:
    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active
    ensure_sheet_headers(ws)
    questions: List[Question] = []
    for r in range(2, ws.max_row + 1):
        q = excel_row_to_question(ws, r)
        if q is not None:
            questions.append(q)
    return questions


def validate_questions(questions: List[Question], require_ids: bool = True) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []
    seen_ids: set[str] = set()
    for idx, q in enumerate(questions, start=1):
        label = q.qid or f"row#{idx}"
        if require_ids:
            if not q.qid:
                errors.append(f"{label}: missing id")
            elif q.qid in seen_ids:
                errors.append(f"{label}: duplicate id '{q.qid}'")
            else:
                seen_ids.add(q.qid)

        if not q.stem:
            errors.append(f"{label}: missing stem")
        if len(q.options) < 2:
            errors.append(f"{label}: fewer than 2 options")
        if len(q.options) > 5:
            errors.append(f"{label}: more than 5 options")

        try:
            ans = normalize_answer(q.answer)
        except Exception as exc:
            errors.append(f"{label}: {exc}")
            ans = "A"
        if q.options and OPTION_LETTERS.index(ans) >= len(q.options):
            errors.append(f"{label}: answer {ans} points to missing option")

        if q.shuffle_mode not in VALID_SHUFFLE_MODES:
            errors.append(f"{label}: invalid shuffle_mode '{q.shuffle_mode}'")

        try:
            if int(q.points) <= 0:
                errors.append(f"{label}: points must be positive")
        except Exception:
            errors.append(f"{label}: invalid points '{q.points}'")

        lowered = [o.casefold().strip() for o in q.options]
        if len(lowered) != len(set(lowered)):
            warnings.append(f"{label}: duplicate option text within question")
        if q.shuffle_mode == "random":
            for opt in q.options:
                if any(p.search(opt) for p in LETTER_REFERENCE_PATTERNS):
                    warnings.append(
                        f"{label}: possible letter-referential option while shuffle_mode=random -> '{opt}'"
                    )
                    break
        if not q.chapter_section:
            warnings.append(f"{label}: missing chapter_section")
    return errors, warnings


def canonical_yaml_from_questions(questions: List[Question]) -> Dict[str, Any]:
    return {
        "defaults": {"shuffle_mode": DEFAULT_SHUFFLE_MODE, "points": DEFAULT_POINTS},
        "questions": [q.to_canonical_yaml_obj(include_id=True) for q in questions],
    }


def randomize_questions_for_export(
    questions: List[Question],
    seed: Optional[int] = None,
    shuffle_questions: bool = True,
    shuffle_options: bool = True,
) -> Tuple[List[Question], Dict[str, Any]]:
    rng = random.Random(seed)
    # deep-ish copy preserving primitive fields
    qlist: List[Question] = [Question(**{**q.__dict__, "options": list(q.options)}) for q in questions]

    if shuffle_questions:
        rng.shuffle(qlist)

    option_permutations: Dict[str, Dict[str, Any]] = {}
    for q in qlist:
        if not shuffle_options or q.shuffle_mode != "random":
            option_permutations[q.qid or q.stem[:20]] = {
                "original_answer": q.answer,
                "new_answer": q.answer,
                "permutation": list(range(len(q.options))),
            }
            continue
        perm = list(range(len(q.options)))
        rng.shuffle(perm)
        old_options = list(q.options)
        old_correct_idx = q.correct_index()
        q.options = [old_options[i] for i in perm]
        new_correct_idx = perm.index(old_correct_idx)
        q.answer = OPTION_LETTERS[new_correct_idx]
        option_permutations[q.qid or q.stem[:20]] = {
            "original_answer": OPTION_LETTERS[old_correct_idx],
            "new_answer": q.answer,
            "permutation": perm,
        }

    manifest = {
        "seed": seed,
        "shuffle_questions": shuffle_questions,
        "shuffle_options": shuffle_options,
        "question_order": [q.qid for q in qlist],
        "option_permutations": option_permutations,
    }
    return qlist, manifest


def questions_to_pyexam_yaml(
    questions: List[Question],
    exam_name: str = DEFAULT_PYEXAM_NAME,
    header: Optional[str] = None,
    include_points: bool = True,
    include_ids_in_text: bool = False,
    include_solution_notes_in_text: bool = False,
) -> Dict[str, Any]:
    payload: Dict[str, Any] = {"name": exam_name, "questions": []}
    if header:
        payload["header"] = header

    pyqs = []
    for q in questions:
        text = q.stem
        if include_ids_in_text and q.qid:
            text = f"[{q.qid}] {text}"
        if include_solution_notes_in_text and q.explanation:
            # Plain text note; pyexam will show this in both versions unless using a custom template.
            text = f"{text}\\n\\n\\emph{{Note: {q.explanation}}}"
        item: Dict[str, Any] = {
            "type": "multiple-choice",
            "text": text,
            "options": [],
        }
        if include_points and q.points:
            item["points"] = int(q.points)
        correct_idx = q.correct_index()
        for i, opt in enumerate(q.options):
            opt_obj: Dict[str, Any] = {"text": opt}
            if i == correct_idx:
                # pyexam examples use "yes" but a YAML boolean should also work.
                opt_obj["correct"] = True
            item["options"].append(opt_obj)
        pyqs.append(item)
    payload["questions"] = pyqs
    return payload


def sanitize_filename_part(s: str) -> str:
    s = (s or "").strip().lower()
    s = re.sub(r"[^a-z0-9._-]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s or "output"
